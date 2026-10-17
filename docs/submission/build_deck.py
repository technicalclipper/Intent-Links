"""Fill the AI Builder Cup submission template with IntentLink's content.

Keeps the template's slides, header strip, fonts and closing slide; drops the
instruction slide, the two optional slides and the spare blank one.
"""
import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

SRC = "template.pptx"
OUT = "IntentLink - AI Builder Cup submission.pptx"

SANS = "Google Sans Flex"
SANS_B = "Google Sans Flex SemiBold"
MONO = "Roboto Mono"

INK = RGBColor(0x0D, 0x0F, 0x13)
CARD = RGBColor(0x16, 0x19, 0x1F)
LINE_D = RGBColor(0x2A, 0x2F, 0x38)
TXT = RGBColor(0x1E, 0x22, 0x28)
DIM = RGBColor(0x5B, 0x62, 0x70)
SOFT = RGBColor(0xC3, 0xC8, 0xD1)
PALE = RGBColor(0xE8, 0xEA, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TINT = RGBColor(0xF4, 0xF5, 0xF7)
BORDER = RGBColor(0xDD, 0xE0, 0xE5)
GREEN = RGBColor(0x2D, 0xD4, 0xA7)
GREEN_D = RGBColor(0x0B, 0x8F, 0x6E)
AMBER = RGBColor(0xF5, 0xB5, 0x44)
RED = RGBColor(0xFF, 0x5C, 0x6C)
RED_D = RGBColor(0xD6, 0x33, 0x47)

prs = Presentation(SRC)
S = list(prs.slides)  # template order, 0-based

CONTENT_TOP = Inches(1.18)
LEFT = Inches(0.42)
WIDTH = Inches(9.16)


# ------------------------------------------------------------------ helpers
def _font(run, size, bold=False, color=TXT, face=None, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color
    name = face or (SANS_B if bold else SANS)
    f.name = name
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs", "a:sym"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", name)
    if bold and name == SANS_B:
        f.bold = False  # the SemiBold face carries the weight


def text(slide, x, y, w, h, paras, name, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: list of paragraphs; each a list of (text, size, opts) runs, or a dict
    {"runs": [...], "space": pt_after, "align": ...}."""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = anchor
    for i, para in enumerate(paras):
        if isinstance(para, list):
            para = {"runs": para}
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        if "space" in para:
            p.space_after = Pt(para["space"])
        if "line" in para:
            p.line_spacing = para["line"]
        for t, size, opts in para["runs"]:
            r = p.add_run()
            r.text = t
            _font(r, size, **opts)
    return tb


def _plain(shape):
    """python-pptx autoshapes reference the theme's effect style, which this
    template renders as a drop shadow; remove the reference, keep our fills."""
    st = shape._element.find(qn("p:style"))
    if st is not None:
        shape._element.remove(st)


def box(slide, x, y, w, h, fill, name, line=None, radius=0.08, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.name = name
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    _plain(s)
    s.text_frame.text = ""
    return s


def dot(slide, cx, cy, d, fill, name, ring=None):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, cx - d // 2, cy - d // 2, d, d)
    s.name = name
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if ring is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = ring
        s.line.width = Pt(1)
    _plain(s)
    return s


def label_in(shape, t, size, color, bold=True, face=None):
    tf = shape.text_frame
    for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, side, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    _font(r, size, bold=bold, color=color, face=face)


def arrow(slide, x1, y1, x2, y2, color, name, width=1.25, head=True, dash=False):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    c.name = name
    c.line.color.rgb = color
    c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    if dash:
        etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
    if head:
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
    return c


def heading(slide, title):
    """The template's question box becomes the slide heading, one line."""
    for shp in slide.shapes:
        if shp.has_text_frame and shp.text_frame.text.strip():
            shp.top = Inches(0.66)
            shp.left = Emu(265350)
            shp.height = Inches(0.45)
            tf = shp.text_frame
            p0 = tf.paragraphs[0]
            # keep the template's first run styling, replace text, drop the rest
            runs = p0.runs
            runs[0].text = title
            for r in runs[1:]:
                r._r.getparent().remove(r._r)
            for br in p0._p.findall(qn("a:br")):
                p0._p.remove(br)
            _font(runs[0], 18, bold=True, color=TXT)
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            return shp
    raise RuntimeError("no heading box")


def chip(slide, x, y, verdict, verb, name, dark=True):
    """ALLOW / DENY chip, the product's rule notation."""
    w_badge = Inches(0.62)
    w = w_badge + Inches(0.12) + Inches(0.095) * len(verb) + Inches(0.16)
    h = Inches(0.32)
    bg = box(slide, x, y, w, h, CARD if dark else WHITE, name, line=LINE_D if dark else BORDER, radius=0.18)
    badge_fill = PALE if verdict == "ALLOW" else None
    b = box(slide, x + Inches(0.05), y + Inches(0.05), w_badge - Inches(0.06), h - Inches(0.1), badge_fill,
            name + " badge", line=None if verdict == "ALLOW" else (PALE if dark else TXT), radius=0.15)
    label_in(b, verdict, 7.5, INK if verdict == "ALLOW" else (PALE if dark else TXT), face=MONO)
    text(slide, x + w_badge + Inches(0.06), y, w - w_badge, h,
         [[(verb, 9.5, {"color": PALE if dark else TXT, "face": MONO})]], name + " verb", anchor=MSO_ANCHOR.MIDDLE)
    return w


def picture(slide, path, x, y, w=None, h=None, name="", frame=True):
    if frame:
        pad = Inches(0.03)
        box(slide, x - pad, y - pad, (w or 0) + 2 * pad, (h or 0) + 2 * pad, None, name + " frame", line=BORDER, radius=0.04)
    p = slide.shapes.add_picture(path, x, y, w, h)
    p.name = name
    return p


# ------------------------------------------------------------ 1 · team details
s = S[0]
tb = next(sh for sh in s.shapes if sh.has_text_frame and "Team Details" in sh.text_frame.text)
values = {"Team name: ": "AP-ICE", "Team leader name: ": "ADAIKAL SUSAN BESCHI J I",
          "Problem Statement: ": "Future of Work & Enterprise Productivity"}
for p in tb.text_frame.paragraphs:
    if p.runs and p.runs[0].text in values:
        lab = p.runs[0]
        _font(lab, 18, bold=True, color=TXT)
        r = copy.deepcopy(lab._r)
        lab._r.addnext(r)
        from pptx.text.text import _Run
        run = _Run(r, p)
        run.text = values[lab.text]
        _font(run, 18, color=TXT)

# --------------------------------------------------------- 3 · brief about it
s = S[2]
heading(s, "Brief about the idea")
text(s, LEFT, CONTENT_TOP, Inches(4.55), Inches(3.9), [
    {"runs": [("Do a task once.", 26, {"bold": True}), ("\n", 26, {}), ("Send it as a link.", 26, {"bold": True, "color": DIM})], "space": 14, "line": 1.0},
    {"runs": [("IntentLink is a Chrome extension that records a task you do in any web app and turns it into a shareable link. "
               "Whoever opens the link has their own AI agent do the same task in their own browser, signed in as them.", 12.5, {})], "space": 10, "line": 1.15},
    {"runs": [("The creator writes the limits in plain English. Gemini turns them into rules, and a policy engine outside the model checks "
               "every action before it happens.", 12.5, {})], "space": 14, "line": 1.15},
    {"runs": [("Don’t send instructions. Send the intent.", 13, {"italic": True, "color": GREEN_D, "bold": True})]},
], "brief text")

cx, cy, cw, ch = Inches(5.25), CONTENT_TOP, Inches(4.33), Inches(3.95)
box(s, cx, cy, cw, ch, INK, "brief card", radius=0.05)
steps = [("01", "RECORD", "Do the task once, in any web app", WHITE, None),
         ("02", "BOUND", "“don’t delete anything”", WHITE, ("DENY", "delete_*")),
         ("03", "SHARE & RUN", "Their agent runs it, inside your rules", GREEN, None)]
node_x = cx + Inches(0.45)
ys = [cy + Inches(0.75), cy + Inches(1.95), cy + Inches(3.15)]
arrow(s, node_x, ys[0], node_x, ys[2], LINE_D, "brief chain", width=1.5, head=False)
for (n, k, desc, col, rule), y in zip(steps, ys):
    d = dot(s, node_x, y, Inches(0.36), col, f"brief node {n}")
    label_in(d, n, 8, INK, face=MONO)
    text(s, node_x + Inches(0.38), y - Inches(0.26), Inches(3.3), Inches(0.24),
         [[(k, 8.5, {"color": SOFT, "face": MONO})]], f"brief step {n} key")
    text(s, node_x + Inches(0.38), y - Inches(0.02), Inches(3.3), Inches(0.3),
         [[(desc, 12, {"color": WHITE, "italic": rule is not None})]], f"brief step {n} text")
    if rule:
        chip(s, node_x + Inches(0.38), y + Inches(0.33), rule[0], rule[1], f"brief rule {n}")
# the boundary, the product's mark
box(s, cx + cw - Inches(0.32), cy + Inches(0.4), Inches(0.05), ch - Inches(0.8), RED, "brief boundary", radius=0.5)

# ---------------------------------------------------------- 4 · opportunities
s = S[3]
heading(s, "Opportunities")
cols = [
    ("How is it different?", [
        "Macro recorders replay clicks and break when a page changes.",
        "AI browser agents adapt, but nothing limits what they do.",
        "IntentLink adapts like an agent and stays bounded like a recording.",
    ]),
    ("How does it solve the problem?", [
        "No more 14-step docs, screenshots or screen-shares.",
        "The creator shows the task once; the receiver presses Run.",
        "Their agent does it in their own browser, with their own accounts.",
    ]),
    ("USP", [
        "The recording is the grant. English is the restriction.",
        "Rules are enforced by code outside the model: the agent cannot talk its way past them.",
        "Every action needs a single-use token, and every block quotes the creator’s own words.",
    ]),
]
gap = Inches(0.22)
cw = (WIDTH - 2 * gap) // 3
for i, (q, pts) in enumerate(cols):
    x = LEFT + i * (cw + gap)
    dark = i == 2
    box(s, x, CONTENT_TOP, cw, Inches(3.0), INK if dark else TINT, f"opp card {i+1}", radius=0.05)
    d = dot(s, x + Inches(0.36), CONTENT_TOP + Inches(0.38), Inches(0.3), GREEN if dark else INK, f"opp icon {i+1}")
    label_in(d, str(i + 1), 9, INK if dark else WHITE, face=MONO)
    text(s, x + Inches(0.6), CONTENT_TOP + Inches(0.26), cw - Inches(0.75), Inches(0.3),
         [[(q, 13, {"bold": True, "color": WHITE if dark else TXT})]], f"opp q {i+1}", anchor=MSO_ANCHOR.MIDDLE)
    paras = []
    for j, t in enumerate(pts):
        emph = dark and j == 0
        paras.append({"runs": [(t, 13 if emph else 11.5, {"color": (GREEN if emph else (SOFT if dark else TXT)), "bold": emph, "italic": emph})],
                      "space": 9, "line": 1.12})
    text(s, x + Inches(0.3), CONTENT_TOP + Inches(0.82), cw - Inches(0.55), Inches(2.1), paras, f"opp body {i+1}")
who_y = CONTENT_TOP + Inches(3.2)
box(s, LEFT, who_y, WIDTH, Inches(0.75), WHITE, "who strip", line=BORDER, radius=0.15)
text(s, LEFT + Inches(0.25), who_y, WIDTH - Inches(0.5), Inches(0.75), [[
    ("Who it’s for  ", 11, {"bold": True}),
    ("teams onboarding new people · support hand-offs · IT rollouts · anyone who writes “how to” docs", 11, {"color": DIM}),
]], "who strip text", anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------------- 5 · features
s = S[4]
heading(s, "Features offered by the solution")
feats = [
    ("Record in any web app", "Chrome extension captures each click by role and label, never by screen position.", INK),
    ("Boundaries in plain English", "“Don’t delete anything” becomes DENY delete_*, compiled by Gemini.", INK),
    ("A consent page per link", "Shows what it will do, what is prohibited, and what the receiver can change.", INK),
    ("Agent in the receiver’s browser", "Google ADK + Gemini resolves each step on their page, as them.", INK),
    ("Policy engine, deny by default", "Allow, ask-first and deny rules checked in code before every action.", GREEN_D),
    ("Single-use action tokens", "Bound to the exact step and element; the extension cannot act without one.", GREEN_D),
    ("Blocked, in your own words", "When the agent reaches the line, it stops and quotes the creator’s sentence.", RED_D),
    ("Parameters & cross-app runs", "Change a recipient or a theme per run; one link can span several apps.", INK),
]
cols, rows = 4, 2
gx, gy = Inches(0.18), Inches(0.2)
fw = (WIDTH - (cols - 1) * gx) // cols
fh = (Inches(3.95) - gy) // rows
for i, (t, d, col) in enumerate(feats):
    r, c = divmod(i, cols)
    x = LEFT + c * (fw + gx)
    y = CONTENT_TOP + r * (fh + gy)
    box(s, x, y, fw, fh, WHITE, f"feature {i+1}", line=BORDER, radius=0.06)
    dt = dot(s, x + Inches(0.32), y + Inches(0.36), Inches(0.28), col, f"feature {i+1} node")
    label_in(dt, f"{i+1:02d}", 7, WHITE, face=MONO)
    text(s, x + Inches(0.2), y + Inches(0.66), fw - Inches(0.36), Inches(0.5),
         [[(t, 12, {"bold": True})]], f"feature {i+1} title")
    text(s, x + Inches(0.2), y + Inches(1.12), fw - Inches(0.36), fh - Inches(1.2),
         [{"runs": [(d, 10.5, {"color": DIM})], "line": 1.1}], f"feature {i+1} text")

# --------------------------------------------------------- 6 · process flow
s = S[5]
heading(s, "Process flow")
lane_h = Inches(1.62)
lanes = [("CREATOR", CONTENT_TOP), ("RECEIVER", CONTENT_TOP + lane_h + Inches(0.38))]
for nm, y in lanes:
    box(s, LEFT, y, WIDTH, lane_h, TINT, f"lane {nm}", radius=0.06)
    text(s, LEFT + Inches(0.18), y + Inches(0.12), Inches(1.5), Inches(0.22),
         [[(nm, 8.5, {"color": DIM, "face": MONO, "bold": False})]], f"lane {nm} label")

def step_box(x, y, w, title, sub, name, fill=INK, tcol=WHITE, scol=SOFT):
    b = box(s, x, y, w, Inches(0.92), fill, name, radius=0.1)
    text(s, x + Inches(0.14), y + Inches(0.13), w - Inches(0.28), Inches(0.3), [[(title, 11.5, {"bold": True, "color": tcol})]], name + " title")
    text(s, x + Inches(0.14), y + Inches(0.45), w - Inches(0.28), Inches(0.42), [{"runs": [(sub, 9, {"color": scol})], "line": 1.05}], name + " sub")
    return b

cy1 = lanes[0][1] + Inches(0.48)
bgap = Inches(0.28)
bw = (WIDTH - Inches(0.4) - 3 * bgap) // 4
xs = [LEFT + Inches(0.2) + i * (bw + bgap) for i in range(4)]
creator = [("Record the task", "Click through it once, in any web app"),
           ("Write the limits", "“don’t delete anything, ask before sending”"),
           ("Gemini compiles", "the steps into a workflow and the sentence into rules"),
           ("Share the link", "published with a consent page, on Cloud Run")]
for i, (t, sub) in enumerate(creator):
    step_box(xs[i], cy1, bw, t, sub, f"creator step {i+1}")
    if i < 3:
        arrow(s, xs[i] + bw, cy1 + Inches(0.46), xs[i + 1], cy1 + Inches(0.46), DIM, f"creator arrow {i+1}")

cy2 = lanes[1][1] + Inches(0.48)
receiver = [("Open & consent", "read what it does and what it may never do, press Run"),
            ("Agent proposes", "the next action on their page (Google ADK + Gemini)"),
            ("Policy engine decides", "in code, against the creator’s rules, every action")]
for i, (t, sub) in enumerate(receiver):
    step_box(xs[i], cy2, bw, t, sub, f"receiver step {i+1}")
    if i < 2:
        arrow(s, xs[i] + bw, cy2 + Inches(0.46), xs[i + 1], cy2 + Inches(0.46), DIM, f"receiver arrow {i+1}")
# the link travels from creator to receiver
arrow(s, xs[3] + bw // 2, cy1 + Inches(0.92), xs[0] + bw // 2, cy2, DIM, "link travels", dash=True)
# two outcomes
ox, ow, oh = xs[3], bw, Inches(0.42)
ok = box(s, ox, cy2 - Inches(0.02), ow, oh, GREEN, "outcome allow", radius=0.2)
text(s, ox + Inches(0.12), cy2 - Inches(0.02), ow - Inches(0.24), oh,
     [[("ALLOW ", 9, {"face": MONO, "bold": False, "color": INK}), ("token → extension acts", 9.5, {"color": INK, "bold": True})]],
     "outcome allow text", anchor=MSO_ANCHOR.MIDDLE)
no = box(s, ox, cy2 + Inches(0.52), ow, oh, RED, "outcome deny", radius=0.2)
text(s, ox + Inches(0.12), cy2 + Inches(0.52), ow - Inches(0.24), oh,
     [[("DENY ", 9, {"face": MONO, "bold": False, "color": INK}), ("blocked, rule quoted", 9.5, {"color": INK, "bold": True})]],
     "outcome deny text", anchor=MSO_ANCHOR.MIDDLE)
arrow(s, xs[2] + bw, cy2 + Inches(0.3), ox, cy2 + Inches(0.19), GREEN_D, "to allow")
arrow(s, xs[2] + bw, cy2 + Inches(0.62), ox, cy2 + Inches(0.73), RED_D, "to deny")

# --------------------------------------------------------- 8 · architecture
s = S[7]
heading(s, "Architecture")
top = CONTENT_TOP
colw = Inches(2.62)
gapx = (WIDTH - 3 * colw) // 2
cols_x = [LEFT + i * (colw + gapx) for i in range(3)]
blocks = [
    ("RECEIVER’S BROWSER", "Chrome extension", ["Recorder: role + label, no coordinates", "Executor: acts only with a token", "Side panel: live run, verdicts"], INK),
    ("CLOUD RUN · NODE / TS", "Policy service", ["Publish: Gemini compiles intent + rules", "Decide & redeem single-use tokens", "Consent, landing, hash-chained audit"], INK),
    ("CLOUD RUN · PYTHON", "Agent (Google ADK)", ["Gemini Flash-Lite picks each action", "before_tool gate asks the policy service", "Holds no authority of its own"], INK),
]
bh = Inches(2.05)
for i, (k, t, lines, fill) in enumerate(blocks):
    x = cols_x[i]
    box(s, x, top, colw, bh, fill, f"arch {t}", radius=0.06)
    text(s, x + Inches(0.2), top + Inches(0.18), colw - Inches(0.4), Inches(0.22), [[(k, 8, {"face": MONO, "color": SOFT})]], f"arch {t} key")
    text(s, x + Inches(0.2), top + Inches(0.42), colw - Inches(0.4), Inches(0.34), [[(t, 14, {"bold": True, "color": WHITE})]], f"arch {t} title")
    text(s, x + Inches(0.2), top + Inches(0.86), colw - Inches(0.4), Inches(1.15),
         [{"runs": [("•  " + ln, 10, {"color": PALE})], "space": 5} for ln in lines], f"arch {t} lines")
mid = top + bh // 2
arrow(s, cols_x[0] + colw, mid - Inches(0.18), cols_x[1], mid - Inches(0.18), GREEN_D, "ext to policy")
text(s, cols_x[0] + colw + Inches(0.02), mid - Inches(0.62), gapx - Inches(0.04), Inches(0.4),
     [{"runs": [("publish", 7.5, {"face": MONO, "color": DIM})], "align": PP_ALIGN.CENTER}, {"runs": [("redeem", 7.5, {"face": MONO, "color": DIM})], "align": PP_ALIGN.CENTER}], "label ext policy")
arrow(s, cols_x[2], mid + Inches(0.18), cols_x[1] + colw, mid + Inches(0.18), GREEN_D, "agent to policy")
text(s, cols_x[1] + colw + Inches(0.02), mid + Inches(0.24), gapx - Inches(0.04), Inches(0.28),
     [{"runs": [("decide", 7.5, {"face": MONO, "color": DIM})], "align": PP_ALIGN.CENTER}], "label agent policy")
# extension <-> agent runs underneath the policy column
yb = top + bh + Inches(0.2)
arrow(s, cols_x[0] + colw // 2, top + bh, cols_x[0] + colw // 2, yb + Inches(0.12), DIM, "ext down", head=False)
arrow(s, cols_x[0] + colw // 2, yb + Inches(0.12), cols_x[2] + colw // 2, yb + Inches(0.12), DIM, "ext to agent", head=False)
arrow(s, cols_x[2] + colw // 2, yb + Inches(0.12), cols_x[2] + colw // 2, top + bh, DIM, "agent up")
text(s, cols_x[1], yb - Inches(0.06), colw, Inches(0.2),
     [{"runs": [("page observation → next action", 7.5, {"face": MONO, "color": DIM})], "align": PP_ALIGN.CENTER}], "label ext agent")
# platform row
py = yb + Inches(0.42)
plats = [("Gemini API", "compile + runtime models"), ("Firestore", "links, policies, audit log"),
         ("Secret Manager", "signing secret, API key"), ("Cloud Build", "images from source")]
pw = (WIDTH - 3 * Inches(0.16)) // 4
for i, (t, sub) in enumerate(plats):
    x = LEFT + i * (pw + Inches(0.16))
    box(s, x, py, pw, Inches(0.62), TINT, f"platform {t}", radius=0.12)
    text(s, x + Inches(0.16), py + Inches(0.08), pw - Inches(0.3), Inches(0.26), [[(t, 11, {"bold": True})]], f"platform {t} title")
    text(s, x + Inches(0.16), py + Inches(0.34), pw - Inches(0.3), Inches(0.22), [[(sub, 9, {"color": DIM})]], f"platform {t} sub")
text(s, LEFT, py + Inches(0.74), WIDTH, Inches(0.3),
     [[("The agent reasons but can’t authorize. The extension acts but can’t authorize. Authority lives in the policy service.", 10.5,
        {"italic": True, "color": TXT})]], "arch thesis")

# ------------------------------------------------------- 9 · technologies
s = S[8]
heading(s, "Technologies used")
groups = [
    ("Google AI", GREEN_D, ["Gemini 3.8 Flash · compiles recordings and rules", "Gemini 3.5 Flash-Lite · runtime agent", "Google ADK 2.11 · agent + tool gate"]),
    ("Google Cloud", INK, ["Cloud Run · policy service + agent", "Firestore · links, policies, audit", "Secret Manager · Cloud Build"]),
    ("Product", INK, ["Chrome extension (Manifest V3, TypeScript)", "Node.js 26 + TypeScript · Python 3.14 + FastAPI", "Zod schemas · signed tokens · hash-chained audit"]),
]
gh = (Inches(3.95) - 2 * Inches(0.18)) // 3
for i, (g, col, items) in enumerate(groups):
    y = CONTENT_TOP + i * (gh + Inches(0.18))
    box(s, LEFT, y, WIDTH, gh, TINT, f"tech row {g}", radius=0.08)
    d = dot(s, LEFT + Inches(0.4), y + gh // 2, Inches(0.3), col, f"tech node {g}")
    text(s, LEFT + Inches(0.72), y, Inches(1.6), gh, [[(g, 14, {"bold": True})]], f"tech group {g}", anchor=MSO_ANCHOR.MIDDLE)
    iw = (WIDTH - Inches(2.45) - Inches(0.2) - 2 * Inches(0.12)) // 3
    for j, it in enumerate(items):
        x = LEFT + Inches(2.45) + j * (iw + Inches(0.12))
        box(s, x, y + Inches(0.17), iw, gh - Inches(0.34), WHITE, f"tech {g} {j+1}", line=BORDER, radius=0.14)
        name, _, sub = it.partition(" · ")
        paras = [[(name, 10.5, {"bold": True})]]
        if sub:
            paras.append([(sub, 9, {"color": DIM})])
        text(s, x + Inches(0.14), y + Inches(0.17), iw - Inches(0.28), gh - Inches(0.34), paras, f"tech {g} {j+1} text", anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------- 11 · snapshots
s = S[10]
heading(s, "Snapshots of the prototype")
iw = Inches(3.62)
ih = Emu(int(iw * 900 / 1440))
py = CONTENT_TOP + Inches(0.08)
picture(s, "shot-landing.png", LEFT + Inches(0.04), py, iw, ih, "snapshot landing")
picture(s, "shot-consent.png", LEFT + iw + Inches(0.32), py, iw, ih, "snapshot consent")
pw = Inches(1.3)
ph = Emu(int(pw * 760 / 380))
picture(s, "shot-panel.png", LEFT + 2 * iw + Inches(0.6), py, pw, ph, "snapshot panel")
cap_y = py + max(ih, ph) + Inches(0.16)
caps = [(LEFT, iw, "Landing page", "Live on Cloud Run, with the extension download and a demo link."),
        (LEFT + iw + Inches(0.29), iw, "Consent page", "What the link will do, what it may never do, and what you can change."),
        (LEFT + 2 * iw + Inches(0.57), pw + Inches(0.05), "Side panel", "Record, set limits, watch the run.")]
for x, w, t, d in caps:
    text(s, x, cap_y, w, Inches(0.7), [[(t, 11.5, {"bold": True})], [(d, 9.5, {"color": DIM})]], f"caption {t}")

# ------------------------------------------------------ 12 · performance
s = S[11]
heading(s, "Prototype performance")
stats = [("226", "automated tests passing", "across schema, policy engine, compiler, API, extension and agent"),
         ("~1 s", "per agent step", "Gemini 3.5 Flash-Lite; steps it can match with certainty skip the model entirely"),
         ("660×", "faster element matching", "600 controls: 79.8 s → 121 ms after a one-pass rewrite"),
         ("6 / 6", "compile judgments correct", "Gemini 3.8 Flash on our golden set, vs 3 / 6 for a Pro model")]
sw = (WIDTH - 3 * Inches(0.18)) // 4
for i, (big, lab, sub) in enumerate(stats):
    x = LEFT + i * (sw + Inches(0.18))
    dark = i == 0
    box(s, x, CONTENT_TOP, sw, Inches(2.75), INK if dark else TINT, f"stat {i+1}", radius=0.06)
    text(s, x + Inches(0.22), CONTENT_TOP + Inches(0.3), sw - Inches(0.4), Inches(0.9),
         [[(big, 34, {"bold": True, "color": GREEN if dark else TXT})]], f"stat {i+1} value")
    text(s, x + Inches(0.22), CONTENT_TOP + Inches(1.2), sw - Inches(0.4), Inches(0.5),
         [[(lab, 12, {"bold": True, "color": WHITE if dark else TXT})]], f"stat {i+1} label")
    text(s, x + Inches(0.22), CONTENT_TOP + Inches(1.68), sw - Inches(0.4), Inches(1.0),
         [{"runs": [(sub, 9.5, {"color": SOFT if dark else DIM})], "line": 1.1}], f"stat {i+1} sub")
box(s, LEFT, CONTENT_TOP + Inches(2.95), WIDTH, Inches(0.95), WHITE, "perf note box", line=BORDER, radius=0.1)
text(s, LEFT + Inches(0.25), CONTENT_TOP + Inches(2.95), WIDTH - Inches(0.5), Inches(0.95), [
    [("Verified live: ", 11, {"bold": True}), ("a YouTube theme change and a Gmail compose-and-send, each recorded once and run end to end in a receiver’s browser. "
      "Creating a link takes 1.5–3 minutes, once, while Gemini compiles the recording and the rules.", 11, {"color": TXT})],
], "perf note", anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------------------ 13 · future
s = S[12]
heading(s, "Future development")
future = [("Chrome Web Store listing", "one-click install instead of developer mode"),
          ("Approvals where people work", "“ask me before sending” routed to Slack or email"),
          ("Team libraries", "versioned links, owners, and who-ran-what across an organisation"),
          ("Admin guardrails", "org-wide rules that every link inherits, on top of the creator’s"),
          ("Scale-out runs", "run state in Firestore so the agent scales horizontally"),
          ("Narrated recordings", "speak while recording; the reasons become part of the intent")]
fw2 = (WIDTH - Inches(0.3)) // 2
for i, (t, d) in enumerate(future):
    c, r = divmod(i, 3)
    x = LEFT + c * (fw2 + Inches(0.3))
    y = CONTENT_TOP + r * Inches(1.32)
    box(s, x, y, fw2, Inches(1.12), TINT, f"future {i+1}", radius=0.1)
    dt = dot(s, x + Inches(0.38), y + Inches(0.56), Inches(0.3), INK if i else GREEN_D, f"future {i+1} node")
    label_in(dt, f"{i+1}", 9, WHITE, face=MONO)
    text(s, x + Inches(0.72), y + Inches(0.2), fw2 - Inches(0.9), Inches(0.75),
         [[(t, 13, {"bold": True})], [(d, 10.5, {"color": DIM})]], f"future {i+1} text", anchor=MSO_ANCHOR.MIDDLE)

# --------------------------------------------------------- 14 · links
s = S[13]
GH = "https://github.com/technicalclipper/Intent-Links"
LIVE = "https://intentlink-api-938318821120.asia-south1.run.app"
tb = next(sh for sh in s.shapes if sh.has_text_frame and "Provide links" in sh.text_frame.text)
links = {"GitHub Public Repository": GH, "Demo Video Link (3 Minutes)": None, "Final Product Link": LIVE}
for p in tb.text_frame.paragraphs:
    key = p.text.strip()
    if key in links and p.runs:
        lab = p.runs[0]
        lab.text = key + ": "
        _font(lab, 16, bold=True, color=TXT)
        from pptx.text.text import _Run
        r = copy.deepcopy(lab._r)
        lab._r.addnext(r)
        run = _Run(r, p)
        url = links[key]
        run.text = url or "to be added after recording"
        _font(run, 14, color=(TXT if url else DIM), italic=url is None)
        if url:
            run.hyperlink.address = url

hy = Inches(2.55)
box(s, LEFT, hy, WIDTH, Inches(1.25), INK, "judge hint", radius=0.08)
dot(s, LEFT + Inches(0.45), hy + Inches(0.62), Inches(0.3), GREEN, "judge hint node")
text(s, LEFT + Inches(0.8), hy, WIDTH - Inches(1.1), Inches(1.25), [
    {"runs": [("Try it in 2 minutes", 15, {"bold": True, "color": WHITE})], "space": 4},
    {"runs": [("Open the product link, download the extension, load it in Chrome, then open the YouTube demo link on the "
               "page and press Run. Your own agent switches YouTube’s theme, and can do nothing else.", 11, {"color": SOFT})], "line": 1.12},
], "judge hint text", anchor=MSO_ANCHOR.MIDDLE)

# ------------------------------------------- drop the slides we don't use
drop = {1, 6, 9, 14}  # instructions, wireframes (optional), cost (optional), spare blank
ids = prs.slides._sldIdLst
for idx in sorted(drop, reverse=True):
    sld = list(ids)[idx]
    prs.part.drop_rel(sld.get(qn("r:id")))
    ids.remove(sld)

prs.save(OUT)
print("saved", OUT)
